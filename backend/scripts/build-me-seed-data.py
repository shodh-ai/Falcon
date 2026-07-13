#!/usr/bin/env python3
"""Parse B.Tech (ME) Updated Time Table.pdf + faculty TT files into me-seed-data.json."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pdfplumber
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
PDF = ROOT / "B.Tech (ME) Updated Time Table.pdf"
OUT = Path(__file__).resolve().parent / "me-seed-data.json"

FACULTY_EMAIL = {
    "NEERAJ": "neeraj.kumar1@mygyanvihar.com",
    "AMIT": "amit.tiwari@mygyanvihar.com",
    "HIMANSHU": "himanshu.vasnani@mygyanvihar.com",
    "RAJ": "raj.kumar@mygyanvihar.com",
    "HOD": "neeraj.kumar1@mygyanvihar.com",
}

FACULTY_NAME_MAP = {
    "neeraj kumar": "NEERAJ",
    "amit tiwari": "AMIT",
    "himanshu vasnani": "HIMANSHU",
    "raj kumar": "RAJ",
}

DAYS = {
    "monday": 1,
    "tuesday": 2,
    "wednesday": 3,
    "thursday": 4,
    "friday": 5,
    "saturday": 6,
}

SEM_PAGES = [
    ("III", "A", "LT-14", 0),
    ("V", "A", "ME Computer Lab", 1),
    ("VII", "A", "Old Computer Lab", 2),
]

TIME_COLS = [
    (1, "09:00", "10:00"),
    (2, "10:00", "10:50"),
    (3, "10:50", "11:40"),
    (4, "11:40", "12:30"),
    (6, "13:30", "14:20"),
    (7, "14:20", "15:10"),
    (8, "15:10", "16:00"),
]

SKIP = re.compile(
    r"(weekly test|library|personal tutor|gate class|nptel|community club|"
    r"kaerb|hcnul|weekly activity|lab projects|tute\b|employability|industrial training seminar|"
    r"election and electoral|field project|field visit|programing software lab|seminar \(sm|"
    r"major project|lab projects preprations|preprations|workshop technology)",
    re.I,
)

CODE_RE = re.compile(
    r"\(?(ME\s*\d[\dA-Z]*|UC\s*[\dA-Z]+|EM\s*\d+|PE\s*\d+|SM\s*\d+|PT\s*\d+|UCFV[\-\dA-Z]*|AE\s*\d+|DME\s*\d+|DHS\s*\d+)\)?",
    re.I,
)


def norm_code(raw: str) -> str:
    code = re.sub(r"\s+", "", raw.upper())
    code = code.replace("UCFV-1", "UCFV1")
    return code


def faculty_key_from_name(name: str) -> str:
    clean = re.sub(r"^dr\.?\s*", "", name.strip(), flags=re.I).lower()
    clean = re.sub(r"\s+", " ", clean)
    for key, val in FACULTY_NAME_MAP.items():
        if key in clean:
            return val
    return "HOD"


def extract_code_and_name(cell: str) -> tuple[str, str] | None:
    if not cell or not isinstance(cell, str):
        return None
    text = " ".join(cell.split())
    if SKIP.search(text):
        return None
    m = CODE_RE.search(text)
    if not m:
        return None
    code = norm_code(m.group(1))
    name = CODE_RE.sub("", text).strip(" -()")
    name = re.sub(r"\s+", " ", name).strip()
    if not name:
        name = code
    return code, name


def parse_faculty_table(table: list[list]) -> dict[str, dict]:
    courses: dict[str, dict] = {}
    for row in table[1:]:
        pairs = [(row[0], row[1], row[2]), (row[4] if len(row) > 4 else None, row[5] if len(row) > 5 else None, row[6] if len(row) > 6 else None)]
        for code_cell, name_cell, fac_cell in pairs:
            if not code_cell or not str(code_cell).strip():
                continue
            code = norm_code(str(code_cell))
            name = " ".join(str(name_cell or code).split())
            fac = faculty_key_from_name(str(fac_cell or ""))
            ctype = "LAB" if any(x in name.upper() for x in ("LAB", "PROJECT", "VISIT", "SEMINAR")) else "THEORY"
            credits = 2 if ctype == "LAB" else 4
            courses[code] = {
                "code": code,
                "name": name,
                "credits": credits,
                "type": ctype,
                "faculty": fac,
            }
    return courses


def parse_student_grid(table: list[list], sem_roman: str, room: str, section: str) -> list[dict]:
    slots: list[dict] = []
    for row in table[1:]:
        day_cell = (row[0] or "").strip().lower()
        if day_cell not in DAYS:
            continue
        day = DAYS[day_cell]
        for col_idx, start, end in TIME_COLS:
            if col_idx >= len(row):
                continue
            cell = row[col_idx]
            if not cell:
                continue
            parsed = extract_code_and_name(str(cell))
            if not parsed:
                continue
            code, _ = parsed
            slots.append(
                {
                    "course": code,
                    "day": day,
                    "start": start,
                    "end": end,
                    "room": room,
                    "section": section,
                }
            )
    return slots


def parse_faculty_xlsx(path: Path, faculty_key: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    time_cols: list[tuple[int, str, str]] = []
    header_idx = None
    for i, row in enumerate(rows):
        if row and isinstance(row[0], str) and row[0].strip().lower() == "days":
            header_idx = i
            for col_idx, cell in enumerate(row):
                if col_idx == 0:
                    continue
                if not cell:
                    continue
                flat = str(cell).replace("\n", " ")
                m = re.search(r"(\d{1,2})[:.](\d{2}).*?(\d{1,2})[:.](\d{2})", flat)
                if m:
                    time_cols.append(
                        (
                            col_idx,
                            f"{int(m.group(1)):02d}:{m.group(2)}",
                            f"{int(m.group(3)):02d}:{m.group(4)}",
                        )
                    )
            break
    if header_idx is None:
        return []

    slots: list[dict] = []
    for row in rows[header_idx + 1 :]:
        if not row or not row[0]:
            continue
        day_cell = str(row[0]).strip().lower()
        if day_cell not in DAYS:
            continue
        day = DAYS[day_cell]
        for col_idx, start, end in time_cols:
            if col_idx >= len(row):
                continue
            cell = row[col_idx]
            if not cell:
                continue
            parsed = extract_code_and_name(str(cell))
            if not parsed:
                continue
            code, _ = parsed
            sem = infer_sem_from_code(code)
            slots.append(
                {
                    "course": code,
                    "semester": sem,
                    "day": day,
                    "start": start,
                    "end": end,
                    "faculty": faculty_key,
                    "room": "ME Faculty Room",
                    "section": "A",
                }
            )
    return slots


def parse_faculty_docx(path: Path, faculty_key: str) -> list[dict]:
    doc = Document(path)
    slots: list[dict] = []
    for table in doc.tables:
        rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows]
        if not rows or rows[0][0].lower() not in ("", "days", "time"):
            continue
        if rows[0][0].lower() == "time":
            continue
        header = rows[0]
        time_cols: list[tuple[int, str, str]] = []
        time_row = rows[1] if len(rows) > 1 and "9:" in rows[1][1] else header
        for col_idx in range(1, len(time_row)):
            flat = time_row[col_idx]
            m = re.search(r"(\d{1,2})[:.](\d{2}).*?(\d{1,2})[:.](\d{2})", flat)
            if m:
                time_cols.append(
                    (
                        col_idx,
                        f"{int(m.group(1)):02d}:{m.group(2)}",
                        f"{int(m.group(3)):02d}:{m.group(4)}",
                    )
                )
        start_row = 2 if rows[0][0].lower() == "days" else 1
        for row in rows[start_row:]:
            if not row or not row[0]:
                continue
            day_cell = row[0].strip().lower()
            if day_cell not in DAYS:
                continue
            day = DAYS[day_cell]
            for col_idx, start, end in time_cols:
                if col_idx >= len(row):
                    continue
                cell = row[col_idx]
                if not cell or SKIP.search(cell):
                    continue
                parsed = extract_code_and_name(cell)
                if not parsed:
                    continue
                code, _ = parsed
                sem = infer_sem_from_code(code)
                slots.append(
                    {
                        "course": code,
                        "semester": sem,
                        "day": day,
                        "start": start,
                        "end": end,
                        "faculty": faculty_key,
                        "room": "ME Faculty Room",
                        "section": "A",
                    }
                )
    return slots


def infer_sem_from_code(code: str) -> str:
    c = norm_code(code)
    if c.startswith("ME3") or c.startswith("UC3") or c in {"UCFV1", "UCEEPI"}:
        return "III"
    if c.startswith("ME30") and c[3:4].isdigit():
        n = int(re.sub(r"\D", "", c[2:5] or "0") or 0)
        if n >= 300 and n < 400:
            if "303" in c or "305" in c or "301" in c or "313" in c:
                return "V"
    if any(x in c for x in ("ME303", "ME305", "ME301", "ME313", "ME351", "ME353", "ME355", "EM301", "UC351", "PT303")):
        return "V"
    if any(x in c for x in ("ME405", "ME407", "ME411", "ME453", "ME459", "PE403", "SM401", "UC401", "EM401")):
        return "VII"
    if c.startswith("ME37"):
        return "III"
    if c.startswith("ME3") and not c.startswith("ME37"):
        return "III"
    return "V"


def main() -> None:
    if not PDF.exists():
        raise SystemExit(f"Missing {PDF}")

    catalog: dict[str, dict] = {}
    semesters: dict[str, dict] = {}

    with pdfplumber.open(PDF) as pdf:
        for sem_roman, section, room, page_idx in SEM_PAGES:
            tables = pdf.pages[page_idx].extract_tables()
            grid = tables[0]
            fac_table = tables[1]
            courses = parse_faculty_table(fac_table)
            for code, course in courses.items():
                course["semester"] = sem_roman
                catalog[code] = course
            slots = parse_student_grid(grid, sem_roman, room, section)
            for slot in slots:
                if slot["course"] in catalog:
                    slot["faculty"] = catalog[slot["course"]]["faculty"]
                else:
                    slot["faculty"] = "HOD"
            semesters[sem_roman] = {
                "room": room,
                "section": section,
                "courses": list(courses.values()),
                "slots": slots,
            }

    faculty_slots: list[dict] = []
    faculty_slots += parse_faculty_xlsx(
        ROOT / "Dr. Himanshu Vasnani INDIVIDUAL TIME TABLE.xlsx", "HIMANSHU"
    )
    faculty_slots += parse_faculty_xlsx(
        ROOT / "Dr. Neeraj Kumar INDIVIDUAL TIME TABLE.xlsx", "NEERAJ"
    )
    faculty_slots += parse_faculty_docx(
        ROOT / "Time Table Amit Tiwari 06 July 2026.docx", "AMIT"
    )
    faculty_slots += parse_faculty_docx(
        ROOT / "Dr. Raj Kumar Time table final.docx", "RAJ"
    )

    for slot in faculty_slots:
        if slot["course"] not in catalog:
            catalog[slot["course"]] = {
                "code": slot["course"],
                "name": slot["course"],
                "credits": 4,
                "type": "THEORY",
                "faculty": slot["faculty"],
                "semester": slot["semester"],
            }

    students = [
        {"email": "anshuman.2549873@mygyanvihar.com", "semester": 3, "section": "A", "name": "Anshuman Singh"},
        {"email": "jalaj.2550454@mygyanvihar.com", "semester": 3, "section": "A", "name": "Jalaj Bansal"},
        {"email": "sunil.2455672@mygyanvihar.com", "semester": 5, "section": "A", "name": "Sunil Kumar"},
        {"email": "raviraj.2455903@mygyanvihar.com", "semester": 5, "section": "A", "name": "Ravi Raj"},
        {"email": "yash.23180717@mygyanvihar.com", "semester": 7, "section": "A", "name": "Yash Singh"},
        {"email": "ravi.2345541@mygyanvihar.com", "semester": 7, "section": "A", "name": "Ravi Kumar"},
    ]

    payload = {
        "academic_year": "2026-2027",
        "program_name": "B.Tech ME",
        "program_code": "BTECH-ME",
        "dept_name": "Mech Engg",
        "faculty": FACULTY_EMAIL,
        "hod_email": FACULTY_EMAIL["NEERAJ"],
        "mentorship_by_semester": {"3": "NEERAJ", "5": "AMIT", "7": "RAJ"},
        "semesters": semesters,
        "faculty_timetable_slots": faculty_slots,
        "catalog": list(catalog.values()),
        "students": students,
        "student_users_to_ensure": students,
    }

    OUT.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"Wrote {OUT}")
    print(
        f"  catalog={len(catalog)} "
        f"sem3_slots={len(semesters['III']['slots'])} "
        f"sem5_slots={len(semesters['V']['slots'])} "
        f"sem7_slots={len(semesters['VII']['slots'])} "
        f"faculty_slots={len(faculty_slots)}"
    )


if __name__ == "__main__":
    main()
